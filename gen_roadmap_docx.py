from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

title = doc.add_heading('Estuda.AI — Roadmap de Melhorias e Correcoes', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run(
    'Levantamento: 22/05/2026  |  Disponibilidade: 15h/dia  |  Total restante: ~93h = 6,2 dias\n'
    'Estrategia: construir valor primeiro, monetizar por ultimo.'
)
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

rule_p = doc.add_paragraph()
rule_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
rule_run = rule_p.add_run(
    'Regra: cada item inicia em branch dedicada. '
    'Ao termino: atualizar README.md > merge na main > push para o GitHub.'
)
rule_run.italic = True
rule_run.font.size = Pt(8.5)
rule_run.font.color.rgb = RGBColor(0x33, 0x33, 0x99)

doc.add_paragraph()


def shade_cell(cell, fill_hex):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def add_fase_heading(doc, text, color_rgb):
    h = doc.add_heading(text, level=1)
    for run in h.runs:
        run.font.color.rgb = RGBColor(*color_rgb)
        run.font.size = Pt(12)


def add_table(doc, headers, rows, col_widths, done_col=3):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = h
        p = cell.paragraphs[0]
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(cell, '1a1a2e')
        p.runs[0].font.color.rgb = RGBColor(0xff, 0xff, 0xff)
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        is_done = done_col < len(row_data) and str(row_data[done_col]).strip().lower() == 'feito'
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            cell.text = str(val)
            run = cell.paragraphs[0].runs[0]
            run.font.size = Pt(7.5) if ci == len(headers) - 1 else Pt(8)
            if is_done:
                run.font.color.rgb = RGBColor(0xaa, 0xaa, 0xaa)
                run.font.italic = True
        if ri % 2 == 0 and not is_done:
            for ci in range(len(headers)):
                shade_cell(row.cells[ci], 'f0f4ff')
        if is_done:
            for ci in range(len(headers)):
                shade_cell(row.cells[ci], 'f5f5f5')
    for i, w in enumerate(col_widths):
        for cell in table.columns[i].cells:
            cell.width = Inches(w)


# FASE 1
add_fase_heading(doc,
    'FASE 1 - Fundacao  (bugs visiveis + qualidade de codigo)  |  10h restantes = 0,7 dias',
    (0x2c, 0x3e, 0x50))
p = doc.add_paragraph('Resolver o que quebra hoje e criar a base que sustenta tudo que vem depois.')
p.runs[0].font.size = Pt(8); p.runs[0].italic = True

add_table(doc,
    ['#', 'Atividade', 'Horas', 'Status', 'Notas'],
    [
        ['1', 'Fix bug progId is not defined ao duplicar aba de Materias', '1h', 'pendente',
         'switchMatTab() referencia progId como variavel livre em onclick'],
        ['2', 'Unificar os dois rows de botoes de programa no painel Materias', '1h', 'pendente',
         'Remover duplicidade de renderizacao'],
        ['3', 'Testes automatizados basicos (tests.js)', '8h', 'pendente',
         'Smoke tests para funcoes criticas; 4.300+ linhas sem teste e risco alto'],
        ['4', 'Otimizar max_tokens por tipo de chamada a API', '2h', 'feito',
         'callClaude() agora aceita maxTokens por feature'],
        ['5', 'Atualizar README.md', '1h', 'feito',
         'README reflete v1.4: SM-2, gamificacao, freemium, multiplos programas'],
    ],
    [0.22, 2.9, 0.55, 0.65, 2.6],
    done_col=3
)

doc.add_paragraph()

# FASE 2
add_fase_heading(doc,
    'FASE 2 - Features diferenciadoras  |  21h = 1,4 dias',
    (0x27, 0x6e, 0x48))
p = doc.add_paragraph('Aumentam engajamento e criam o valor percebido que justificara a cobranca futura.')
p.runs[0].font.size = Pt(8); p.runs[0].italic = True

add_table(doc,
    ['#', 'Atividade', 'Horas', 'Status', 'Notas'],
    [
        ['6', 'Diagnostico de lacunas de conhecimento', '8h', 'pendente',
         'Identifica topicos com mais erros no simulado > sugere historinha ou revisao SM-2'],
        ['7', 'Compartilhamento de programas de estudo', '6h', 'pendente',
         'Export/import de programa com materias e topicos (JSON)'],
        ['8', 'Estimativa de prontidao para a prova', '4h', 'pendente',
         '"Com seu ritmo atual, cobrira X% do conteudo ate a data da prova"'],
        ['9', 'Suporte a LaTeX/MathJax nas respostas', '3h', 'pendente',
         'Essencial para Economia (formulas em texto puro degradam o conteudo)'],
    ],
    [0.22, 2.9, 0.55, 0.65, 2.6],
    done_col=3
)

doc.add_paragraph()

# FASE 3
add_fase_heading(doc,
    'FASE 3 - Retencao e produto completo  (roadmap pendente)  |  34h = 2,3 dias',
    (0xe6, 0x7e, 0x22))
p = doc.add_paragraph('Features que aumentam retorno diario e completam o que ja foi prometido.')
p.runs[0].font.size = Pt(8); p.runs[0].italic = True

add_table(doc,
    ['#', 'Atividade', 'Horas', 'Status', 'Notas'],
    [
        ['10', 'PWA: manifest.json + Service Worker offline', '6h', 'pendente',
         'Instalacao na tela inicial do celular; funciona offline real'],
        ['11', 'Push notifications para revisao SM-2', '8h', 'pendente',
         'Lembrete no dia correto; SM-2 perde metade do valor sem o alerta'],
        ['12', 'Onboarding wizard para novos usuarios', '6h', 'pendente',
         'Wizard 3 passos: criar programa > materia > primeira historinha'],
        ['13', 'Historico de notas das redacoes com grafico', '4h', 'pendente',
         'Chart.js ja disponivel; salvar nota por redacao no Supabase'],
        ['14', 'Exportacao de historinhas e flashcards para PDF', '6h', 'pendente',
         'window.print() + CSS de impressao ou jsPDF'],
        ['15', 'Estatisticas por materia e por periodo', '4h', 'pendente',
         'Breakdown de acertos/erros por materia no simulado'],
    ],
    [0.22, 2.9, 0.55, 0.65, 2.6],
    done_col=3
)

doc.add_paragraph()

# FASE 4
add_fase_heading(doc,
    'FASE 4 - Monetizacao  (cobrar por um produto que ja entrega valor)  |  28h = 1,9 dias',
    (0xc0, 0x39, 0x2b))
p = doc.add_paragraph('So faz sentido cobrar depois que o produto estiver solido e com usuarios engajados.')
p.runs[0].font.size = Pt(8); p.runs[0].italic = True

add_table(doc,
    ['#', 'Atividade', 'Horas', 'Status', 'Notas'],
    [
        ['16', 'Integracao Stripe real (substituir ativarPremiumDemo())', '16h', 'pendente',
         'Checkout session, webhook, atualizacao de plano no Supabase'],
        ['17', 'Proxy API via Supabase Edge Function', '8h', 'pendente',
         'Plataforma chama Anthropic com chave propria; fim da dependencia da chave do usuario'],
        ['18', 'Rate limiting real no backend (Supabase)', '4h', 'pendente',
         'Contador de geracoes via RLS/trigger - nao so frontend; impede burla via console'],
    ],
    [0.22, 2.9, 0.55, 0.65, 2.6],
    done_col=3
)

doc.add_paragraph()

# RESUMO
doc.add_heading('Resumo Geral', level=2)
add_table(doc,
    ['Fase', 'Horas restantes', 'Dias (15h/dia)'],
    [
        ['Fase 1 - Fundacao', '10h', '0,7 dias'],
        ['Fase 2 - Features diferenciadoras', '21h', '1,4 dias'],
        ['Fase 3 - Retencao e produto completo', '34h', '2,3 dias'],
        ['Fase 4 - Monetizacao', '28h', '1,9 dias'],
        ['TOTAL RESTANTE', '93h', '~6,2 dias'],
    ],
    [3.5, 1.5, 2.0],
    done_col=99
)

doc.add_paragraph()

# REGRA
doc.add_heading('Regra de execucao de cada item', level=2)
steps = [
    '1.  git checkout -b <tipo>/<descricao-curta>',
    '2.  Implementar',
    '3.  Atualizar README.md com a melhoria/correcao',
    '4.  git commit -m "tipo: descricao"',
    '5.  git checkout main && git merge <branch>',
    '6.  git push origin main',
]
for s in steps:
    p = doc.add_paragraph(s, style='List Bullet')
    p.runs[0].font.size = Pt(9)

doc.add_paragraph()
footer_p = doc.add_paragraph(
    'Atualizado em 22/05/2026 - estrategia invertida: valor antes de monetizacao.'
)
footer_p.runs[0].font.size = Pt(7.5)
footer_p.runs[0].italic = True
footer_p.runs[0].font.color.rgb = RGBColor(0xaa, 0xaa, 0xaa)

doc.save('ROADMAP.docx')
print('ROADMAP.docx gerado com sucesso.')
